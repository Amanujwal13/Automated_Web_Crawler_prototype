from __future__ import annotations

from pathlib import Path
import openpyxl
from exporters import export_all, export_excel, export_summary


def test_export_excel_produces_valid_workbook(tmp_path: Path):
    records = [
        {
            "title": "A Light in the Attic",
            "price": "£51.77",
            "upc": "a897fe39b1053632",
            "availability": "In stock",
            "domain": "books.toscrape.com",
            "source_url": "https://books.toscrape.com/catalogue/1",
            "final_url": "https://books.toscrape.com/catalogue/1",
            "fetched_at": "2026-10-29T10:00:00Z",
            "tables": [
                {
                    "table_id": 1,
                    "caption": "Product Info",
                    "headers": ["Key", "Val"],
                    "rows": [["UPC", "a897fe39b1053632"]],
                }
            ],
        }
    ]

    xlsx_path = export_excel(records, tmp_path)
    assert xlsx_path.exists()

    wb = openpyxl.load_workbook(xlsx_path)
    sheet_names = wb.sheetnames
    assert "Records" in sheet_names
    assert "Extracted Tables" in sheet_names
    assert "Metadata" in sheet_names

    ws_records = wb["Records"]
    assert ws_records.cell(row=2, column=1).value == "A Light in the Attic"


def test_export_summary_produces_readable_report(tmp_path: Path):
    records = [
        {
            "title": "Test Book",
            "domain": "books.toscrape.com",
            "byline": "Author X",
            "tables": [],
        }
    ]

    summary_file = export_summary(records, tmp_path, run_id="20261029-100000-abcd")
    assert summary_file.exists()

    content = summary_file.read_text(encoding="utf-8")
    assert "WEB CRAWLER STUDIO — STAGE 2 RUN SUMMARY REPORT" in content
    assert "20261029-100000-abcd" in content
    assert "books.toscrape.com" in content


def test_export_all_creates_all_artifacts(tmp_path: Path):
    records = [{"title": "Demo Record", "domain": "example.com"}]
    artifacts = export_all(records, tmp_path, run_id="20261029-100000-abcd")

    assert set(artifacts.keys()) == {"json", "csv", "excel", "summary"}
    for path in artifacts.values():
        assert path.exists()
