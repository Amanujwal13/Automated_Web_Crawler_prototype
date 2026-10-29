from __future__ import annotations

import csv
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, List, Optional

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

CSV_DEFAULT_HEADERS = [
    "title",
    "price",
    "upc",
    "availability",
    "byline",
    "published_at",
    "summary",
    "table_count",
    "extraction_mode",
    "domain",
    "source_url",
    "final_url",
    "fetched_at",
]


def export_json(
    records: list[dict[str, Any]],
    output_dir: Path | str,
    overwrite: bool = False,
) -> Path:
    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    target_file = out_dir / "data.json"
    if target_file.exists() and not overwrite:
        raise FileExistsError(f"Export target already exists; refusing to overwrite: {target_file}")

    serialized = json.dumps(records, indent=2, ensure_ascii=False) + "\n"
    target_file.write_text(serialized, encoding="utf-8")
    return target_file


def export_csv(
    records: list[dict[str, Any]],
    output_dir: Path | str,
    overwrite: bool = False,
) -> Path:
    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    target_file = out_dir / "data.csv"
    if target_file.exists() and not overwrite:
        raise FileExistsError(f"Export target already exists; refusing to overwrite: {target_file}")

    # Build union of headers while prioritizing standard fields
    fieldnames = list(CSV_DEFAULT_HEADERS)
    for rec in records:
        for k in rec.keys():
            if k not in fieldnames and k != "tables":  # nested tables excluded from flat CSV
                fieldnames.append(k)

    with target_file.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for rec in records:
            # Flatten any dict/list values for CSV cells
            flattened = {}
            for k, v in rec.items():
                if k == "tables":
                    continue
                if isinstance(v, (list, dict)):
                    flattened[k] = json.dumps(v, ensure_ascii=False)
                else:
                    flattened[k] = v
            writer.writerow(flattened)

    return target_file


def export_excel(
    records: list[dict[str, Any]],
    output_dir: Path | str,
    overwrite: bool = False,
) -> Path:
    """
    Stage 2 Excel exporter using openpyxl.
    Generates data.xlsx with styled 'Records', 'Extracted Tables', and 'Metadata' sheets.
    """
    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    target_file = out_dir / "data.xlsx"
    if target_file.exists() and not overwrite:
        raise FileExistsError(f"Export target already exists; refusing to overwrite: {target_file}")

    if not OPENPYXL_AVAILABLE:
        # Fallback: create empty or notice file if openpyxl is unavailable
        target_file.write_bytes(b"")
        return target_file

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4C2365", end_color="4C2365", fill_type="solid")
    table_header_fill = PatternFill(start_color="7F5D99", end_color="7F5D99", fill_type="solid")
    align_left = Alignment(horizontal="left", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")

    # Sheet 1: Records
    ws_records = wb.create_sheet(title="Records")
    fieldnames = list(CSV_DEFAULT_HEADERS)
    for rec in records:
        for k in rec.keys():
            if k not in fieldnames and k != "tables":
                fieldnames.append(k)

    ws_records.append(fieldnames)
    for col_idx in range(1, len(fieldnames) + 1):
        cell = ws_records.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center

    for rec in records:
        row_vals = []
        for field in fieldnames:
            val = rec.get(field)
            if isinstance(val, (list, dict)):
                val = json.dumps(val, ensure_ascii=False)
            row_vals.append(val if val is not None else "")
        ws_records.append(row_vals)

    # Sheet 2: Extracted Tables
    ws_tables = wb.create_sheet(title="Extracted Tables")
    table_cols = ["Source URL", "Domain", "Table #", "Caption", "Headers", "Row #", "Row Data"]
    ws_tables.append(table_cols)
    for col_idx in range(1, len(table_cols) + 1):
        cell = ws_tables.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = table_header_fill
        cell.alignment = align_center

    for rec in records:
        page_url = rec.get("final_url") or rec.get("source_url") or ""
        domain = rec.get("domain") or ""
        tables = rec.get("tables", [])
        for tbl in tables:
            caption = tbl.get("caption", "Table")
            t_id = tbl.get("table_id", 1)
            headers_str = " | ".join(tbl.get("headers", []))
            for r_idx, row in enumerate(tbl.get("rows", [])):
                row_str = " | ".join(str(c) for c in row)
                ws_tables.append([page_url, domain, t_id, caption, headers_str, r_idx + 1, row_str])

    # Sheet 3: Metadata
    ws_meta = wb.create_sheet(title="Metadata")
    ws_meta.append(["Property", "Value"])
    ws_meta.cell(row=1, column=1).font = header_font
    ws_meta.cell(row=1, column=1).fill = header_fill
    ws_meta.cell(row=1, column=2).font = header_font
    ws_meta.cell(row=1, column=2).fill = header_fill

    domains = sorted(list(set(r.get("domain", "") for r in records if r.get("domain"))))
    modes = sorted(list(set(r.get("extraction_mode", "profile") for r in records)))
    metadata_entries = [
        ("Total Records Crawled", len(records)),
        ("Target Domains", ", ".join(domains) if domains else "None"),
        ("Extraction Modes", ", ".join(modes)),
        ("Export Timestamp (UTC)", datetime.now(timezone.utc).isoformat()),
        ("Engine", "Web Crawler Studio Stage 2"),
    ]
    for prop, val in metadata_entries:
        ws_meta.append([prop, val])

    wb.save(target_file)
    return target_file


def export_summary(
    records: list[dict[str, Any]],
    output_dir: Path | str,
    run_id: str,
    stats: Optional[dict[str, Any]] = None,
    overwrite: bool = False,
) -> Path:
    """
    Writes summary.txt in output/<run_id>/ detailing metrics and results.
    """
    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    target_file = out_dir / "summary.txt"
    if target_file.exists() and not overwrite:
        raise FileExistsError(f"Export target already exists; refusing to overwrite: {target_file}")

    domains = sorted(list(set(r.get("domain", "") for r in records if r.get("domain"))))
    total_tables = sum(len(r.get("tables", [])) for r in records)
    total_with_author = sum(1 for r in records if r.get("byline") or r.get("author"))

    summary_content = (
        f"===========================================================\n"
        f" WEB CRAWLER STUDIO — STAGE 2 RUN SUMMARY REPORT\n"
        f"===========================================================\n"
        f"Run ID:            {run_id}\n"
        f"Completed At:      {datetime.now(timezone.utc).isoformat()}\n"
        f"Total Records:     {len(records)}\n"
        f"Target Domains:    {', '.join(domains) if domains else 'None'}\n"
        f"Extracted Tables:  {total_tables}\n"
        f"Records w/ Author: {total_with_author}\n"
        f"Export Directory:  {out_dir}\n"
        f"Export Artifacts:  data.json, data.csv, data.xlsx, summary.txt\n"
        f"===========================================================\n"
    )

    target_file.write_text(summary_content, encoding="utf-8")
    return target_file


def export_all(
    records: list[dict[str, Any]],
    output_dir: Path | str,
    run_id: str,
    overwrite: bool = False,
) -> dict[str, Path]:
    """
    Master export function producing all formats in output/<run_id>/.
    """
    return {
        "json": export_json(records, output_dir, overwrite=overwrite),
        "csv": export_csv(records, output_dir, overwrite=overwrite),
        "excel": export_excel(records, output_dir, overwrite=overwrite),
        "summary": export_summary(records, output_dir, run_id, overwrite=overwrite),
    }
